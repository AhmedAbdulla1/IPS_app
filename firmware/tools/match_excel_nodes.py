import openpyxl
import requests

url = 'https://dqnmxlljqiqgqmntzvcx.supabase.co/rest/v1/nodes?select=*'
headers = {
    'apikey': 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImRxbm14bGxqcWlxZ3FtbnR6dmN4Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODYyNjQwMzIsImV4cCI6MjEwMTg0MDAzMn0.RBcHQxjlZy3Cv-57gaRwM_BISoYaFJfTfJs9kxTLT7w',
    'Authorization': 'Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImRxbm14bGxqcWlxZ3FtbnR6dmN4Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODYyNjQwMzIsImV4cCI6MjEwMTg0MDAzMn0.RBcHQxjlZy3Cv-57gaRwM_BISoYaFJfTfJs9kxTLT7w'
}

r = requests.get(url, headers=headers)
nodes_list = r.json()
nodes = {n['node_id']: n for n in nodes_list}

wb = openpyxl.load_workbook('firmware/node_name.xlsx')
sheet = wb.active

print(f"Excel sheet title: {sheet.title}, total rows: {sheet.max_row}")
print("-" * 80)

for r in range(1, sheet.max_row + 1):
    name = sheet.cell(r, 1).value
    pt_num = sheet.cell(r, 2).value
    x_val = sheet.cell(r, 3).value
    y_val = sheet.cell(r, 4).value
    note = sheet.cell(r, 5).value

    if pt_num is None:
        if name or note:
            print(f"Row {r:2d}: [Non-data row] col1={name}, col2={pt_num}, col3={x_val}, col4={y_val}, col5={note}")
        continue

    try:
        pt_int = int(pt_num)
        supabase_id = 400 + pt_int
        node = nodes.get(supabase_id)
        if node:
            print(f"Row {r:2d}: Pt {pt_int:2d} -> Node {supabase_id}: (x={node['x']}, y={node['y']}) | Excel Name: '{name}' | Supabase Name: '{node['name_ar']}'")
        else:
            print(f"Row {r:2d}: Pt {pt_int:2d} -> Node {supabase_id}: *** NOT FOUND IN SUPABASE *** | Excel Name: '{name}'")
    except Exception as e:
        print(f"Row {r:2d}: Pt value '{pt_num}' could not be converted: {e} | Name: '{name}'")

print("-" * 80)
print(f"All Supabase node IDs: {sorted(nodes.keys())}")
