import openpyxl
import requests

url = 'https://dqnmxlljqiqgqmntzvcx.supabase.co/rest/v1/nodes?select=*'
headers = {
    'apikey': 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImRxbm14bGxqcWlxZ3FtbnR6dmN4Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODYyNjQwMzIsImV4cCI6MjEwMTg0MDAzMn0.RBcHQxjlZy3Cv-57gaRwM_BISoYaFJfTfJs9kxTLT7w',
    'Authorization': 'Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImRxbm14bGxqcWlxZ3FtbnR6dmN4Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODYyNjQwMzIsImV4cCI6MjEwMTg0MDAzMn0.RBcHQxjlZy3Cv-57gaRwM_BISoYaFJfTfJs9kxTLT7w'
}

r = requests.get(url, headers=headers)
nodes = {n['node_id']: n for n in r.json()}

excel_path = 'firmware/node_name.xlsx'
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

updated_rows = []
not_found_rows = []

for r in range(4, sheet.max_row + 1):
    name = sheet.cell(r, 1).value
    pt_val = sheet.cell(r, 2).value
    if pt_val is not None:
        try:
            pt_int = int(pt_val)
            node_id = 400 + pt_int
            if node_id in nodes:
                n = nodes[node_id]
                sheet.cell(r, 3, value=n['x'])
                sheet.cell(r, 4, value=n['y'])
                updated_rows.append((r, name, pt_int, node_id, n['x'], n['y'], n.get('name_ar')))
            else:
                not_found_rows.append((r, name, pt_int, node_id))
        except Exception:
            pass

wb.save(excel_path)

print(f"Successfully updated {len(updated_rows)} rows in '{excel_path}':")
for r, name, pt, node_id, x, y, sup_name in updated_rows:
    print(f"  Row {r:2d}: '{name}' (Pt {pt}) -> Node {node_id}: x={x}, y={y} (Supabase: '{sup_name}')")

if not_found_rows:
    print(f"\nNodes not found in Supabase ({len(not_found_rows)}):")
    for r, name, pt, node_id in not_found_rows:
        print(f"  Row {r:2d}: '{name}' (Pt {pt}) -> Looked for Node {node_id}")
